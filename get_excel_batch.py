import openpyxl
import sys
import json
import os

BASE_DIR = r"c:\Users\ENRIQ\OneDrive\Documents\PROYECTO CON MONICA"
EXCEL_PATH = os.path.join(BASE_DIR, "BASE_DE_DATOS_DEPURADA_WHATSAPP_2026.xlsx")
PROCESSED_FILE = os.path.join(BASE_DIR, "contactos_procesados.json")

def is_valid_phone(phone):
    if not phone or len(phone) != 10 or not phone.isdigit():
        return False
    # No puede empezar con 0 ni 1
    if phone[0] in ['0', '1']:
        return False
    # No puede tener 4 o más dígitos idénticos consecutivos (ej: 7777, 0000, 1111)
    import re
    if re.search(r'(\d)\1{3,}', phone):
        return False
    # Debe tener al menos 4 dígitos únicos distintos
    if len(set(phone)) < 4:
        return False
    # Descartar secuencias comunes
    if phone in ['1234567890', '0123456789', '9876543210', '0000000000', '1111111111']:
        return False
    return True

def get_batch(limit=30):
    processed_phones = set()
    if os.path.exists(PROCESSED_FILE):
        try:
            with open(PROCESSED_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
                processed_phones = set(data.keys())
        except Exception:
            pass

    if not os.path.exists(EXCEL_PATH):
        return []

    wb = None
    temp_copy = None
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
    except PermissionError:
        import subprocess
        import tempfile
        temp_dir = tempfile.gettempdir()
        temp_copy = os.path.join(temp_dir, f"batch_excel_copy_{os.getpid()}.xlsx")
        cmd = f"Copy-Item -Path '{EXCEL_PATH}' -Destination '{temp_copy}' -Force"
        subprocess.run(["powershell", "-Command", cmd], capture_output=True)
        if os.path.exists(temp_copy):
            wb = openpyxl.load_workbook(temp_copy, read_only=True)
        else:
            return []
    except Exception as e:
        return []

    ws = wb['Contactos WhatsApp (40k)']

    batch = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or len(row) < 4:
            continue
        
        phone = str(row[3]).strip() if row[3] is not None else ''
        if not is_valid_phone(phone) or phone in processed_phones:
            continue
            
        nombre = str(row[1] or 'Contacto').strip()
        primer_nombre = str(row[2] or 'Amigo(a)').strip()
        colonia = str(row[6] or 'Navojoa').strip()
        municipio = str(row[7] or 'Navojoa').strip()
        lada = str(row[8] or '642').strip()
        
        batch.append({
            'phone': phone,
            'nombre': nombre,
            'primer_nombre': primer_nombre,
            'colonia': colonia,
            'municipio': municipio,
            'lada': lada
        })
        
        if len(batch) >= limit:
            break

    try:
        wb.close()
    except Exception:
        pass

    if temp_copy and os.path.exists(temp_copy):
        try:
            os.remove(temp_copy)
        except Exception:
            pass

    return batch

if __name__ == "__main__":
    limit = 30
    for arg in sys.argv[1:]:
        if arg.startswith('--limit='):
            try:
                limit = int(arg.split('=')[1])
            except:
                limit = 30
    
    batch = get_batch(limit)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(batch, ensure_ascii=False))
